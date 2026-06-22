#!/usr/bin/env python3
"""
Generate PowerPoint presentation from YAML data.

Usage:
    python generate_slides.py

Requirements:
    pip install python-pptx pyyaml
"""

import yaml
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor


def load_presentation_data(yaml_file='presentation-data.yaml'):
    """Load presentation data from YAML file."""
    with open(yaml_file, 'r') as f:
        return yaml.safe_load(f)


def create_title_slide(prs, slide_data):
    """Create title slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[0])  # Title slide layout
    title = slide.shapes.title
    subtitle = slide.placeholders[1]

    title.text = slide_data['title']

    # Combine all text sections for subtitle
    subtitle_text = []
    for section in slide_data.get('sections', []):
        if section['type'] == 'text':
            subtitle_text.append(section['content'])

    subtitle.text = '\n'.join(subtitle_text)


def create_content_slide(prs, slide_data):
    """Create content slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and content layout
    title = slide.shapes.title
    title.text = slide_data['title']

    # Get content area
    left = Inches(0.5)
    top = Inches(1.5)
    width = Inches(9)
    height = Inches(5)

    current_top = top

    for section in slide_data.get('sections', []):
        section_type = section['type']

        if section_type == 'subtitle':
            # Add subtitle
            txBox = slide.shapes.add_textbox(left, current_top, width, Inches(0.5))
            tf = txBox.text_frame
            tf.text = section['content']
            tf.paragraphs[0].font.size = Pt(18)
            tf.paragraphs[0].font.bold = True
            current_top += Inches(0.6)

        elif section_type == 'text':
            # Add text content
            txBox = slide.shapes.add_textbox(left, current_top, width, Inches(1))
            tf = txBox.text_frame
            tf.text = section['content']
            tf.word_wrap = True
            tf.paragraphs[0].font.size = Pt(14)
            current_top += Inches(0.8)

        elif section_type == 'bullets':
            # Add bulleted list
            if 'title' in section:
                txBox = slide.shapes.add_textbox(left, current_top, width, Inches(0.4))
                tf = txBox.text_frame
                tf.text = section['title']
                tf.paragraphs[0].font.size = Pt(16)
                tf.paragraphs[0].font.bold = True
                current_top += Inches(0.5)

            txBox = slide.shapes.add_textbox(left, current_top, width, Inches(2))
            tf = txBox.text_frame
            for item in section['items']:
                p = tf.add_paragraph()
                p.text = item
                p.level = 0
                p.font.size = Pt(12)
            current_top += Inches(2.2)

        elif section_type == 'table':
            # Add table (simplified - would need proper table creation)
            table_text = []
            if 'title' in section:
                table_text.append(section['title'])

            # Headers
            table_text.append(' | '.join(section.get('headers', [])))
            table_text.append('-' * 50)

            # Rows
            for row in section.get('rows', []):
                table_text.append(' | '.join(str(cell) for cell in row))

            txBox = slide.shapes.add_textbox(left, current_top, width, Inches(2))
            tf = txBox.text_frame
            tf.text = '\n'.join(table_text)
            tf.paragraphs[0].font.size = Pt(10)
            tf.paragraphs[0].font.name = 'Courier New'
            current_top += Inches(2.2)

        elif section_type == 'code':
            # Add code block
            txBox = slide.shapes.add_textbox(left, current_top, width, Inches(3))
            tf = txBox.text_frame
            tf.text = section['content']
            tf.paragraphs[0].font.size = Pt(9)
            tf.paragraphs[0].font.name = 'Courier New'
            current_top += Inches(3.2)

        elif section_type == 'diagram':
            # Add diagram as monospace text
            txBox = slide.shapes.add_textbox(left, current_top, width, Inches(2))
            tf = txBox.text_frame
            tf.text = section['content']
            tf.paragraphs[0].font.size = Pt(10)
            tf.paragraphs[0].font.name = 'Courier New'
            current_top += Inches(2.2)

        elif section_type == 'callout':
            # Add highlighted callout
            txBox = slide.shapes.add_textbox(left, current_top, width, Inches(0.8))
            tf = txBox.text_frame
            tf.text = section['content']
            tf.paragraphs[0].font.size = Pt(14)
            tf.paragraphs[0].font.bold = True
            # Add background color (yellow highlight)
            fill = txBox.fill
            fill.solid()
            fill.fore_color.rgb = RGBColor(255, 255, 200)
            current_top += Inches(1)


def generate_presentation(yaml_file='presentation-data.yaml', output_file='presentation.pptx'):
    """Generate PowerPoint presentation from YAML data."""
    print(f"Loading data from {yaml_file}...")
    data = load_presentation_data(yaml_file)

    print("Creating presentation...")
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # Process each slide
    for slide_data in data['slides']:
        slide_num = slide_data['slide_number']
        print(f"  Creating slide {slide_num}: {slide_data['title']}")

        if slide_num == 1:
            create_title_slide(prs, slide_data)
        else:
            create_content_slide(prs, slide_data)

    # Save presentation
    print(f"Saving presentation to {output_file}...")
    prs.save(output_file)
    print(f"✅ Presentation generated successfully: {output_file}")
    print(f"   Total slides: {len(data['slides'])}")


if __name__ == '__main__':
    try:
        generate_presentation()
    except FileNotFoundError:
        print("❌ Error: presentation-data.yaml not found")
        print("   Make sure you're in the correct directory")
    except Exception as e:
        print(f"❌ Error: {e}")
        import traceback
        traceback.print_exc()
